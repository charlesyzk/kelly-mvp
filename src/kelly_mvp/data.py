"""Strict CSV/Excel ingestion plus an explicitly labelled daily-data fallback."""

from __future__ import annotations

import calendar
import csv
import io
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from math import isfinite
from pathlib import Path
from typing import Iterable
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .config import FREQUENCIES


@dataclass(frozen=True, slots=True)
class PriceRow:
    date: date
    symbol: str
    adjusted_close: float


def _parse_rows(handle: Iterable[str]) -> list[PriceRow]:
    rows: list[PriceRow] = []
    seen: set[tuple[str, date]] = set()
    reader = csv.DictReader(handle)
    required = {"date", "symbol", "adjusted_close"}
    missing = required.difference(reader.fieldnames or ())
    if missing:
        raise ValueError(f"input CSV is missing columns: {sorted(missing)}")
    for line_number, raw in enumerate(reader, start=2):
        try:
            observed = date.fromisoformat(raw["date"].strip())
            symbol = raw["symbol"].strip()
            close = float(raw["adjusted_close"])
        except (AttributeError, TypeError, ValueError) as exc:
            raise ValueError(f"invalid input at line {line_number}") from exc
        if not symbol:
            raise ValueError(f"empty symbol at line {line_number}")
        if not isfinite(close) or close <= 0:
            raise ValueError(f"adjusted_close must be finite and positive at line {line_number}")
        key = (symbol, observed)
        if key in seen:
            raise ValueError(f"duplicate symbol/date at line {line_number}: {symbol} {observed}")
        seen.add(key)
        rows.append(PriceRow(observed, symbol, close))
    if not rows:
        raise ValueError("input CSV contains no data rows")
    return sorted(rows, key=lambda row: (row.symbol, row.date))


def parse_daily_prices(text: str) -> list[PriceRow]:
    if not isinstance(text, str) or not text.strip():
        raise ValueError("input CSV is empty")
    return _parse_rows(io.StringIO(text.lstrip("\ufeff")))


def load_daily_prices(path: str | Path) -> list[PriceRow]:
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(f"input CSV does not exist: {source}")
    return parse_daily_prices(source.read_text(encoding="utf-8-sig"))


def parse_price_workbook(content: bytes) -> dict[str, list[PriceRow]]:
    """Load provider-supplied daily/weekly/monthly sheets from XLSX.

    Each present sheet must be named daily, weekly or monthly and contain
    date, symbol and adjusted_close columns. Frequencies are never inferred
    from filenames or silently manufactured inside this loader.
    """

    if not isinstance(content, bytes) or not content:
        raise ValueError("input workbook is empty")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError) as exc:
        raise ValueError("input workbook is not a valid XLSX file") from exc
    output: dict[str, list[PriceRow]] = {}
    for frequency in FREQUENCIES:
        if frequency not in workbook.sheetnames:
            continue
        sheet = workbook[frequency]
        iterator = sheet.iter_rows(values_only=True)
        try:
            header = next(iterator)
        except StopIteration:
            raise ValueError(f"sheet {frequency} is empty") from None
        names = [str(value).strip().lower() if value is not None else "" for value in header]
        required = ("date", "symbol", "adjusted_close")
        if any(name not in names for name in required):
            raise ValueError(f"sheet {frequency} must contain {required}")
        index = {name: names.index(name) for name in required}
        rows: list[PriceRow] = []
        seen: set[tuple[str, date]] = set()
        for row_number, values in enumerate(iterator, 2):
            if not values or all(value in (None, "") for value in values):
                continue
            raw_date = values[index["date"]]
            observed = raw_date.date() if isinstance(raw_date, datetime) else raw_date
            if not isinstance(observed, date):
                try:
                    observed = date.fromisoformat(str(raw_date).strip())
                except ValueError as exc:
                    raise ValueError(f"invalid date in {frequency}!{row_number}") from exc
            symbol = str(values[index["symbol"]] or "").strip()
            try:
                close = float(values[index["adjusted_close"]])
            except (TypeError, ValueError) as exc:
                raise ValueError(f"invalid adjusted_close in {frequency}!{row_number}") from exc
            if not symbol or not isfinite(close) or close <= 0:
                raise ValueError(f"invalid price row in {frequency}!{row_number}")
            key = (symbol, observed)
            if key in seen:
                raise ValueError(f"duplicate symbol/date in {frequency}!{row_number}")
            seen.add(key)
            rows.append(PriceRow(observed, symbol, close))
        if rows:
            output[frequency] = sorted(rows, key=lambda row: (row.symbol, row.date))
    if not output:
        raise ValueError("workbook has no daily, weekly or monthly price sheets")
    return output


def load_price_workbook(path: str | Path) -> dict[str, list[PriceRow]]:
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(f"input workbook does not exist: {source}")
    return parse_price_workbook(source.read_bytes())


def daily_prices_to_csv(rows: Iterable[PriceRow]) -> str:
    """Serialize normalized prices using the project's public CSV contract."""

    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(("date", "symbol", "adjusted_close"))
    count = 0
    for row in rows:
        writer.writerow((row.date.isoformat(), row.symbol, format(row.adjusted_close, ".15g")))
        count += 1
    if count == 0:
        raise ValueError("cannot serialize an empty price series")
    return output.getvalue()


def _period_key(observed: date, frequency: str) -> tuple[int, int]:
    if frequency == "weekly":
        iso = observed.isocalendar()
        return iso.year, iso.week
    if frequency == "monthly":
        return observed.year, observed.month
    raise ValueError(f"unsupported aggregate frequency: {frequency}")


def aggregate_prices(rows: Iterable[PriceRow], frequency: str) -> list[PriceRow]:
    ordered = sorted(rows, key=lambda row: row.date)
    if frequency == "daily":
        return ordered
    if frequency not in {"weekly", "monthly"}:
        raise ValueError(f"unsupported frequency: {frequency}")
    groups: list[list[PriceRow]] = []
    for row in ordered:
        key = _period_key(row.date, frequency)
        if not groups or _period_key(groups[-1][-1].date, frequency) != key:
            groups.append([])
        groups[-1].append(row)
    # Earlier groups are proven complete by the existence of the next period.
    # The final group is only retained when the input reaches an unambiguous
    # calendar boundary; otherwise it is conservatively excluded.
    complete = groups[:-1]
    if groups:
        final = groups[-1]
        last = final[-1].date
        if frequency == "weekly":
            end = last + timedelta(days=4 - last.weekday())
        else:
            end = date(last.year, last.month, calendar.monthrange(last.year, last.month)[1])
        if last >= end:
            complete.append(final)
    return [group[-1] for group in complete]


def split_by_symbol(rows: Iterable[PriceRow]) -> dict[str, list[PriceRow]]:
    result: dict[str, list[PriceRow]] = {}
    for row in rows:
        result.setdefault(row.symbol, []).append(row)
    return result
