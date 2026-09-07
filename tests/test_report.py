import tempfile
import unittest
from datetime import date, timedelta

from openpyxl import load_workbook

from kelly_mvp import PriceRow, StrategyConfig, run_backtest
from kelly_mvp.report import write_outputs


class ReportTests(unittest.TestCase):
    def test_writes_csv_excel_and_text_outputs(self):
        rows = [PriceRow(date(2020,1,1)+timedelta(days=i), "X", 100*1.001**i) for i in range(18)]
        config = StrategyConfig(
            windows={"daily":5,"weekly":3,"monthly":2},
            minimum_matches={"daily":3,"weekly":2,"monthly":1},
            bootstrap_blocks={"daily":2,"weekly":2,"monthly":1},
            bootstrap_repetitions=20,
        )
        result = run_backtest({name: rows for name in ("daily","weekly","monthly")}, config)
        with tempfile.TemporaryDirectory() as directory:
            paths = write_outputs(result, directory, config)
            self.assertEqual(len(paths), 7)
            self.assertTrue(all(path.is_file() for path in paths))
            workbook = load_workbook(paths[5], read_only=True)
            self.assertEqual(set(workbook.sheetnames), {"summary","statistics","signals","periods","trades"})
            self.assertIn("5% 正式支持", paths[6].read_text(encoding="utf-8"))
            self.assertIn("position_type", paths[2].read_text(encoding="utf-8-sig"))
            self.assertIn("action", paths[3].read_text(encoding="utf-8-sig"))


if __name__ == "__main__":
    unittest.main()
